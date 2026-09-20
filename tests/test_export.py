import openpyxl

from authorfinder.export import write_results_xlsx


def _result(url, **author):
    base = {
        "name": "Jane Doe",
        "profile_url": "https://s.com/author/jane-doe",
        "job_title": "Reporter",
        "bio": "Jane covers policy.",
        "email": "jane@site.com",
        "linkedin": "https://www.linkedin.com/in/jane-doe",
        "twitter": "https://twitter.com/janedoe",
        "social_links": [{"platform": "twitter", "url": "https://twitter.com/janedoe"}],
        "organization": "Example News",
        "location": "Washington",
    }
    base.update(author)
    return {"article_url": url, "author": base, "status": "success"}


def _read(path):
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, row)))
    workbook.close()
    return headers, rows


def test_write_results_xlsx_writes_headers_and_values(tmp_path):
    path = tmp_path / "out.xlsx"
    written = write_results_xlsx([_result("https://s.com/1")], str(path))
    assert written == 1

    headers, rows = _read(path)
    assert "Article URL" in headers
    assert "Email" in headers
    assert "LinkedIn" in headers
    assert "Twitter / X" in headers
    assert rows[0]["Name"] == "Jane Doe"
    assert rows[0]["Email"] == "jane@site.com"
    assert rows[0]["Status"] == "success"
    assert rows[0]["Social Links"] == "twitter: https://twitter.com/janedoe"


def test_write_results_xlsx_freezes_header_and_wraps_text(tmp_path):
    path = tmp_path / "out.xlsx"
    write_results_xlsx([_result("https://s.com/1")], str(path))
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    assert sheet.freeze_panes == "A2"
    assert sheet.cell(row=2, column=1).alignment.wrap_text is True
    workbook.close()


def test_write_results_xlsx_omits_article_urls_unless_deduped(tmp_path):
    plain = tmp_path / "plain.xlsx"
    write_results_xlsx([_result("https://s.com/1")], str(plain))
    headers, _ = _read(plain)
    assert "Article URLs" not in headers

    deduped = tmp_path / "deduped.xlsx"
    row = _result("https://s.com/1")
    row["article_urls"] = ["https://s.com/1", "https://s.com/2"]
    write_results_xlsx([row], str(deduped))
    headers, rows = _read(deduped)
    assert "Article URLs" in headers
    assert rows[0]["Article URLs"] == "https://s.com/1; https://s.com/2"


def test_write_results_xlsx_handles_a_missing_author(tmp_path):
    path = tmp_path / "out.xlsx"
    write_results_xlsx(
        [{"article_url": "https://s.com/1", "status": "blocked", "author": {}}], str(path)
    )
    _, rows = _read(path)
    assert rows[0]["Status"] == "blocked"
    assert rows[0]["Name"] is None
