"""Shared Excel (.xlsx) export for crawl results.

The CLI (`--excel`) and the desktop GUI both write their workbooks through
`write_results_xlsx`, so the column set can never drift between the two.
"""

from __future__ import annotations

# Fields written to Excel, in column order.
EXPORT_FIELDS = (
    "article_url",
    "status",
    "name",
    "profile_url",
    "job_title",
    "bio",
    "email",
    "linkedin",
    "twitter",
    "social_links",
    "organization",
    "location",
)

# Nicer labels than a plain title-case of the field name.
_HEADERS = {
    "article_url": "Article URL",
    "article_urls": "Article URLs",
    "profile_url": "Profile URL",
    "social_links": "Social Links",
    "linkedin": "LinkedIn",
    "twitter": "Twitter / X",
    "job_title": "Job Title",
}


def _header(field: str) -> str:
    return _HEADERS.get(field, field.replace("_", " ").title())


def _columns_and_rows(results: list[dict]) -> tuple[list[str], list[dict]]:
    """Flatten results into spreadsheet columns plus one dict per row."""
    columns = list(EXPORT_FIELDS)
    if any(result.get("article_urls") for result in results):
        columns.insert(1, "article_urls")

    rows = []
    for result in results:
        author = result.get("author") or {}
        authors = result.get("authors") or []
        social = author.get("social_links") or []
        social_str = "; ".join(f"{s['platform']}: {s['url']}" for s in social)
        urls = result.get("article_urls") or [result.get("article_url")]
        if len(authors) > 1:
            name = "; ".join(a.get("name", "") for a in authors if a.get("name"))
            email = "; ".join(a.get("email", "") for a in authors if a.get("email"))
            linkedin = "; ".join(a.get("linkedin", "") for a in authors if a.get("linkedin"))
            twitter = "; ".join(a.get("twitter", "") for a in authors if a.get("twitter"))
            job_title = "; ".join(a.get("job_title", "") for a in authors if a.get("job_title"))
            organization = "; ".join(a.get("organization", "") for a in authors if a.get("organization"))
            profile_url = "; ".join(a.get("profile_url", "") for a in authors if a.get("profile_url"))
            location = "; ".join(a.get("location", "") for a in authors if a.get("location"))
            all_socials = []
            for a in authors:
                all_socials.extend(a.get("social_links", []))
            social_str = "; ".join(f"{s['platform']}: {s['url']}" for s in all_socials)
        else:
            name = author.get("name")
            email = author.get("email")
            linkedin = author.get("linkedin")
            twitter = author.get("twitter")
            job_title = author.get("job_title")
            organization = author.get("organization")
            profile_url = author.get("profile_url")
            location = author.get("location")
        rows.append(
            {
                "article_url": result.get("article_url"),
                "article_urls": "; ".join(url for url in urls if url),
                "status": result.get("status"),
                "name": name,
                "profile_url": profile_url,
                "job_title": job_title,
                "bio": author.get("bio"),
                "email": email,
                "linkedin": linkedin,
                "twitter": twitter,
                "social_links": social_str,
                "organization": organization,
                "location": location,
            }
        )
    return columns, rows


def write_results_xlsx(results: list[dict], path: str) -> int:
    """Write results to an .xlsx workbook. Returns the number of data rows."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise RuntimeError(
            "openpyxl is required for Excel export. Install it: pip install openpyxl"
        ) from exc

    columns, rows = _columns_and_rows(results)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "AuthorFinder Results"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_align = Alignment(wrap_text=True, vertical="top")

    for column_index, field in enumerate(columns, 1):
        cell = sheet.cell(row=1, column=column_index, value=_header(field))
        cell.font = header_font
        cell.fill = header_fill

    for row_index, row in enumerate(rows, 2):
        for column_index, field in enumerate(columns, 1):
            cell = sheet.cell(row=row_index, column=column_index, value=row.get(field))
            cell.alignment = wrap_align

    # Approximate auto-width, capped so the sheet stays readable.
    for column_index, field in enumerate(columns, 1):
        longest = max(
            len(_header(field)),
            max((len(str(row.get(field) or "")) for row in rows), default=0),
        )
        letter = openpyxl.utils.get_column_letter(column_index)
        sheet.column_dimensions[letter].width = min(longest + 4, 50)

    sheet.freeze_panes = "A2"
    workbook.save(path)
    workbook.close()
    return len(rows)
